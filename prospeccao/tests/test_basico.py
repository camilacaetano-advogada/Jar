import random
import sys
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from prospect import caixa, config, contacts, sender, templates
from prospect.store import Store

FUSO = ZoneInfo("America/Sao_Paulo")


@pytest.fixture
def cfg():
    c = config.carregar_config()
    c["envio"]["horario"]["dias"] = [0, 1, 2, 3, 4, 5, 6]
    c["envio"]["horario"]["inicio"] = "00:00"
    c["envio"]["horario"]["fim"] = "23:59"
    return c


@pytest.fixture
def tpl():
    return config.carregar_templates()


def planilha_xlsx(caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Marca", "E-mail de Contato para Parceria", "Ramo / Segmento"])
    ws.append(["Glow Cosméticos", "oi@glow.com.br", "Beleza e cosméticos"])
    ws.append(["Sem Email", "", "Moda"])
    ws.append(["Email Ruim", "isso-nao-e-email", "Moda"])
    ws.append(["Glow Repetida", "OI@glow.com.br", "Beleza"])
    ws.append(["Fit Co", "contato@fit.co; outro@fit.co", "Moda fitness e academia"])
    ws.append(["Festival X", "prod@festivalx.com", "Eventos e shows"])
    wb.save(caminho)


def test_leitura_e_relatorio(tmp_path):
    p = tmp_path / "c.xlsx"
    planilha_xlsx(p)
    lista, rel = contacts.ler_contatos(p)
    assert [c.email for c in lista] == ["oi@glow.com.br", "contato@fit.co", "prod@festivalx.com"]
    assert rel["sem_email"] == 1 and len(rel["invalido"]) == 1 and len(rel["duplicado"]) == 1


def test_csv_ponto_e_virgula(tmp_path):
    p = tmp_path / "c.csv"
    p.write_text("Marca;E-mail de Contato para Parceria;Ramo / Segmento\nLoja A;a@loja.com;Moda\n", encoding="utf-8")
    lista, _ = contacts.ler_contatos(p)
    assert lista[0].marca == "Loja A"


def test_segmentos(tpl):
    assert templates.classificar("Moda fitness e academia", tpl) == "fitness"
    assert templates.classificar("Beleza e cosméticos", tpl) == "beleza"
    assert templates.classificar("Escritório de contabilidade", tpl) == "juridico"
    assert templates.classificar("Eventos e shows", tpl) == "eventos"
    assert templates.classificar("Pet shop", tpl) == "lifestyle"


def test_email_sem_contato_nem_promessa_falsa(cfg, tpl):
    cfg["consumo_geral"] = False
    c = contacts.Contato(marca="Glow", email="oi@glow.com.br", segmento="Beleza")
    m = templates.montar(c, cfg, tpl)
    assert "Oi, pessoal da Glow" in m.texto or "time da Glow" in m.texto
    assert "{" not in m.texto and "}" not in m.texto
    assert "já consumo" not in m.texto  # sem motivo_consumo, não afirma que consome
    assert "@camilaarcanjo.c" in m.texto and "@advocunty" in m.texto
    assert "SAIR" in m.texto and "portfolio-ugc" in m.texto
    cfg["consumo_geral"] = True
    assert "já consumo" in templates.montar(c, cfg, tpl).texto
    cfg["consumo_geral"] = False
    c.motivo_consumo = "Uso o sérum de vocês há 1 ano."
    assert "já consumo" in templates.montar(c, cfg, tpl).texto


def test_followup_menciona_original(cfg, tpl):
    c = contacts.Contato(marca="Glow", email="oi@glow.com.br", segmento="Beleza", contato="Ana Souza")
    m = templates.montar(c, cfg, tpl, etapa=2, assunto_original="Assunto X")
    assert m.assunto == "Re: Assunto X" and "Ana" in m.texto and "creator" in m.texto


def rodar(cfg, tpl, store, lista, enviados, **kw):
    kw.setdefault("dormir", lambda s: None)
    kw.setdefault("mostrar", lambda *a: None)
    kw.setdefault("agora_fn", lambda: datetime.now(FUSO))
    return sender.executar(cfg, tpl, store, lista, real=True, enviar_fn=lambda mime: enviados.append(mime) or mime["Message-ID"],
                           usuario="camila@gmail.com", **kw)


def lista3():
    return [contacts.Contato(marca=f"M{i}", email=f"m{i}@x.com", segmento="Moda") for i in range(3)]


def test_nunca_envia_duas_vezes_e_retoma(cfg, tpl):
    store, enviados = Store(":memory:"), []
    r1 = rodar(cfg, tpl, store, lista3(), enviados, limite=2)
    assert r1.enviados == 2
    r2 = rodar(cfg, tpl, store, lista3(), enviados)
    assert r2.enviados == 1 and len(enviados) == 3
    r3 = rodar(cfg, tpl, store, lista3(), enviados)
    assert r3.enviados == 0 and len(enviados) == 3


def test_limite_diario_e_intervalo(cfg, tpl):
    cfg["envio"]["limite_diario"] = 2
    esperas, store, enviados = [], Store(":memory:"), []
    r = rodar(cfg, tpl, store, lista3(), enviados, dormir=esperas.append, rng=random.Random(1))
    assert r.enviados == 2 and "limite diário" in r.parou_por
    assert len(esperas) == 1 and 60 <= esperas[0] <= 90


def test_bloqueio_e_erro(cfg, tpl):
    store, enviados = Store(":memory:"), []
    store.bloquear("m0@x.com", "teste", "agora")
    r = rodar(cfg, tpl, store, lista3(), enviados)
    assert r.enviados == 2 and r.pulados
    assert all("m0@x.com" not in m["To"] for m in enviados)


def test_followup_so_para_quem_nao_respondeu(cfg, tpl):
    store, enviados = Store(":memory:"), []
    lista = lista3()
    rodar(cfg, tpl, store, lista, enviados)
    antigo = (datetime.now(FUSO) - timedelta(days=7)).isoformat()
    store.db.execute("UPDATE envios SET enviado_em=?", (antigo,))
    store.db.commit()
    store.definir_relacao("m1@x.com", "respondeu", "agora")
    enviados.clear()
    r = rodar(cfg, tpl, store, lista, enviados)
    assert r.enviados == 2
    assert {m["To"] for m in enviados} == {"m0@x.com", "m2@x.com"}
    assert all(m["Subject"].startswith("Re:") and m["In-Reply-To"] for m in enviados)


def test_status_na_planilha(cfg, tpl, tmp_path):
    p = tmp_path / "c.xlsx"
    planilha_xlsx(p)
    lista, _ = contacts.ler_contatos(p)
    store, enviados = Store(":memory:"), []
    rodar(cfg, tpl, store, lista, enviados, planilha=p)
    store.definir_relacao("prod@festivalx.com", "negociacao", "agora")
    sender.sincronizar_planilha(store, p)
    ws = openpyxl.load_workbook(p).active
    cab = [c.value for c in ws[1]]
    st = {r[1].value: r[cab.index("Status")].value for r in ws.iter_rows(min_row=2)}
    assert st["oi@glow.com.br"] == "Enviado" and st["prod@festivalx.com"] == "Em negociação"


def test_pdf_anexado(cfg, tpl, tmp_path):
    pdf = tmp_path / "kit.pdf"
    pdf.write_bytes(b"%PDF-1.4 teste")
    from prospect import mailer

    m = templates.montar(contacts.Contato(marca="A", email="a@a.com"), cfg, tpl)
    mime = mailer.montar_mime(m, "a@a.com", "camila@gmail.com", cfg, pdf)
    assert any(part.get_filename() == "kit.pdf" for part in mime.iter_attachments())
    assert mime.get_body(("html",)) is not None


def test_bounce_e_sair(tpl):
    bounce = b"From: Mail Delivery Subsystem <mailer-daemon@googlemail.com>\nTo: c@gmail.com\nSubject: Delivery failure\n\nA mensagem para m1@x.com falhou."
    resp = b"From: Ana <m2@x.com>\nTo: c@gmail.com\nSubject: Re: oi\n\nSAIR por favor"
    env = {"m1@x.com", "m2@x.com"}
    assert caixa.classificar_mensagem(bounce, env) == [("bounce", "m1@x.com")]
    assert caixa.classificar_mensagem(resp, env) == [("respondeu", "m2@x.com"), ("sair", "m2@x.com")]


def test_janela_de_horario(cfg):
    cfg2 = config.carregar_config()
    seg_10h = datetime(2026, 10, 5, 10, 0, tzinfo=FUSO)
    sab_10h = datetime(2026, 10, 10, 10, 0, tzinfo=FUSO)
    assert sender.dentro_da_janela(cfg2, seg_10h) and not sender.dentro_da_janela(cfg2, sab_10h)
    assert sender.proxima_janela(cfg2, sab_10h).weekday() == 0


class _Exec:
    def __init__(self, v): self.v = v
    def execute(self): return self.v


class _FakeGmail:
    """Imita só o pedaço da API do Gmail que o app usa."""
    def __init__(self, raws=()):
        self.raws, self.enviados = list(raws), []

    def users(self): return self
    def messages(self): return self

    def send(self, userId, body):
        self.enviados.append(body["raw"])
        return _Exec({"id": "1"})

    def list(self, userId, q, pageToken=None):
        return _Exec({"messages": [{"id": str(i)} for i in range(len(self.raws))]})

    def get(self, userId, id, format):
        import base64
        return _Exec({"raw": base64.urlsafe_b64encode(self.raws[int(id)]).decode().rstrip("=")})


def test_enviar_gmail_manda_mime_em_base64(cfg, tpl):
    import base64
    from email import message_from_bytes, policy
    from prospect import mailer

    m = templates.montar(contacts.Contato(marca="A", email="a@a.com"), cfg, tpl)
    mime = mailer.montar_mime(m, "a@a.com", "camila@gmail.com", cfg)
    svc = _FakeGmail()
    assert mailer.enviar_gmail(mime, svc) == mime["Message-ID"]
    enviado = message_from_bytes(base64.urlsafe_b64decode(svc.enviados[0]), policy=policy.default)
    assert enviado["To"] == "a@a.com" and enviado["Subject"] == m.assunto


def test_varrer_caixa_via_api(tpl):
    store = Store(":memory:")
    store.iniciar("m2@x.com", 1, "M2", "a", "2026-10-05")
    store.concluir("m2@x.com", 1, True, "agora")
    resp = b"From: Ana <m2@x.com>\nTo: c@gmail.com\nSubject: Re: oi\n\nSAIR por favor"
    c = caixa.varrer(_FakeGmail([resp]), store, "agora")
    assert c == {"respondeu": 1, "sair": 1, "bounce": 0}
    assert store.bloqueado("m2@x.com") and store.relacao("m2@x.com") == "respondeu"


def test_hospedagem_usa_email_proprio(cfg, tpl):
    c = contacts.Contato(marca="Hotel Sol", email="r@hotelsol.com", segmento="Hotel / Hospedagem", contato="Ana Souza")
    m = templates.montar(c, cfg, tpl)
    assert m.segmento == "hospedagem"
    assert "Poços de Caldas" in m.texto and "Conferência Mineira da Advocacia" in m.texto and "diárias" in m.texto
    assert "SAIR" in m.texto and "Hotel Sol" in m.texto
    assert "já consumo" not in m.texto


def test_hospedagem_sai_primeiro_na_fila(cfg, tpl):
    lista = [
        contacts.Contato(marca="Glow", email="g@glow.com", segmento="Beleza"),
        contacts.Contato(marca="Moda1", email="m@moda.com", segmento="Moda"),
        contacts.Contato(marca="Hotel A", email="a@hotel.com", segmento="Permuta p/ Estadia"),
        contacts.Contato(marca="Hotel B", email="b@hotel.com", segmento="Hotel"),
    ]
    enviados = []
    rodar(cfg, tpl, Store(":memory:"), lista, enviados)
    assert [m["To"] for m in enviados] == ["a@hotel.com", "b@hotel.com", "g@glow.com", "m@moda.com"]
