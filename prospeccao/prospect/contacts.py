"""Leitura da planilha de contatos (CSV/XLSX) e escrita do status de volta nela."""
import csv
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-']+@[A-Za-z0-9\-]+(\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}$")
COLUNAS_STATUS = ["Status", "Enviado em", "Observações"]


def norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(c for c in t if not unicodedata.combining(c)).lower().strip()


@dataclass
class Contato:
    marca: str
    email: str
    segmento: str = ""
    contato: str = ""
    cidade: str = ""
    motivo_consumo: str = ""
    linha: int = 0
    extras: dict = field(default_factory=dict)


class Planilha:
    """Tabela em memória (cabeçalho + linhas) que sabe salvar no formato de origem."""

    def __init__(self, caminho: Path):
        self.caminho = Path(caminho)
        self.xlsx = self.caminho.suffix.lower() in (".xlsx", ".xlsm")
        if self.xlsx:
            import openpyxl

            self._wb = openpyxl.load_workbook(self.caminho)
            ws = self._wb.active
            linhas = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
        else:
            bruto = None
            for enc in ("utf-8-sig", "latin-1"):
                try:
                    bruto = self.caminho.read_text(encoding=enc)
                    self._enc = "utf-8-sig"
                    break
                except UnicodeDecodeError:
                    continue
            try:
                self._delim = csv.Sniffer().sniff(bruto[:4096], delimiters=",;\t").delimiter
            except csv.Error:
                self._delim = ","
            linhas = [[c.strip() for c in r] for r in csv.reader(bruto.splitlines(), delimiter=self._delim)]
        linhas = [r for r in linhas if any(c for c in r)]
        if not linhas:
            raise ValueError(f"Planilha vazia: {self.caminho}")
        self.cabecalho, self.linhas = linhas[0], linhas[1:]

    def indice(self, predicado) -> int | None:
        for i, nome in enumerate(self.cabecalho):
            if predicado(norm(nome)):
                return i
        return None

    def _garantir_coluna(self, nome: str) -> int:
        i = self.indice(lambda n: n == norm(nome))
        if i is None:
            self.cabecalho.append(nome)
            for r in self.linhas:
                r.append("")
            i = len(self.cabecalho) - 1
        return i

    def _col_email(self) -> int | None:
        return self.indice(lambda n: "mail" in n)

    def atualizar(self, status_por_email: dict[str, tuple[str, str, str]]) -> None:
        """status_por_email: email -> (status, enviado_em, observacao)."""
        ie = self._col_email()
        if ie is None:
            return
        cols = [self._garantir_coluna(n) for n in COLUNAS_STATUS]
        for r in self.linhas:
            while len(r) < len(self.cabecalho):
                r.append("")
            achados = [e.lower() for e in re.findall(r"[^\s,;<>]+@[^\s,;<>]+", r[ie])]
            for e in achados:
                if e in status_por_email:
                    for c, v in zip(cols, status_por_email[e]):
                        r[c] = v
                    break
        self.salvar()

    def salvar(self) -> None:
        if self.xlsx:
            ws = self._wb.active
            ws.delete_rows(1, ws.max_row)
            ws.append(self.cabecalho)
            for r in self.linhas:
                ws.append(r)
            self._wb.save(self.caminho)
        else:
            with open(self.caminho, "w", encoding=self._enc, newline="") as f:
                w = csv.writer(f, delimiter=self._delim)
                w.writerow(self.cabecalho)
                w.writerows(self.linhas)


def ler_contatos(caminho: Path) -> tuple[list[Contato], dict]:
    """Devolve contatos válidos + relatório do que foi ignorado e por quê."""
    pl = Planilha(caminho)
    i_email = pl._col_email()
    if i_email is None:
        raise ValueError("Não achei a coluna de e-mail (o título precisa conter 'mail').")
    i_marca = pl.indice(lambda n: n.startswith("marca") or n in ("empresa", "nome_marca", "nome da marca"))
    i_seg = pl.indice(lambda n: "segmento" in n or "ramo" in n)
    i_cid = pl.indice(lambda n: "cidade" in n)
    i_motivo = pl.indice(lambda n: "motivo" in n)
    i_contato = pl.indice(lambda n: ("contato" in n or "responsavel" in n or n == "nome") and "mail" not in n)
    i_marca = 0 if i_marca is None else i_marca

    def pega(r, i):
        return r[i].strip() if i is not None and i < len(r) else ""

    rel = {"total": len(pl.linhas), "sem_email": 0, "invalido": [], "duplicado": [], "vazias_ok": 0}
    vistos, contatos = set(), []
    for n, r in enumerate(pl.linhas, start=2):
        bruto = pega(r, i_email)
        if not bruto:
            rel["sem_email"] += 1
            continue
        candidatos = [e for e in re.split(r"[\s,;/]+", bruto) if "@" in e]
        email = next((e.strip(".<>").lower() for e in candidatos if EMAIL_RE.match(e.strip(".<>"))), None)
        if not email:
            rel["invalido"].append((n, bruto))
            continue
        if email in vistos:
            rel["duplicado"].append((n, email))
            continue
        vistos.add(email)
        contatos.append(
            Contato(
                marca=pega(r, i_marca), email=email, segmento=pega(r, i_seg), contato=pega(r, i_contato),
                cidade=pega(r, i_cid), motivo_consumo=pega(r, i_motivo), linha=n,
            )
        )
    return contatos, rel
